import hashlib
import json
import mimetypes
import re
import unicodedata
import zipfile
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from threading import Lock
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from uygulama.cekirdek.gunlukleyici import gunlukleyici
from uygulama.cekirdek.konfigurasyon import konfigurasyon
from uygulama.cekirdek.veritabani import baglam_yoneticisi
from uygulama.modeller.veritabani_modelleri import tablolari_olustur
from uygulama.servisler.metin_parcalayici import MetinParcalayici

XLSX_MIME_TIPI = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IZINLI_MIME_TIPLERI = {XLSX_MIME_TIPI}
OOXML_IMZASI = b"PK"
AZAMI_SIKISTIRILMAMIS_BOYUT = 250 * 1024 * 1024
AZAMI_ARSIV_GIRISI = 20000
AZAMI_SIKISTIRMA_ORANI = 200
TEHLIKELI_ARSIV_PARCALARI = (
    "vbaproject.bin",
    "/embeddings/",
    "/oleobjects/",
    "/activex/",
    "/externallinks/",
)
FORMUL_BASLANGICI = re.compile(r"^[\s\x00-\x1f]*[=+@]")
_ICE_AKTARIM_KILIDI = Lock()


@dataclass(frozen=True)
class SSSKaydi:
    kategori: str
    soru: str
    cevap: str
    kaynak_dokuman: str | None
    harici_kod: str | None
    modul: str | None
    kullanici_rolu: str | None
    arama_metni: str
    kaynak_sayfa: str
    kaynak_satir: int
    kayit_ozeti: str
    kayit_anahtari: str


@dataclass(frozen=True)
class IceAktarimSonucu:
    surum_numarasi: int
    aktarilan_kayit_sayisi: int
    olusturulan_parca_sayisi: int
    arsivlenen_kayit_sayisi: int
    eklenen_kayit_sayisi: int
    guncellenen_kayit_sayisi: int
    korunan_kayit_sayisi: int
    yinelenen_dosya: bool
    dosya_ozeti: str

    def sozluge_donustur(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class KuruCalistirmaSonucu:
    toplam_kayit_sayisi: int
    eklenen_kayit_sayisi: int
    guncellenen_kayit_sayisi: int
    korunan_kayit_sayisi: int
    arsivlenen_kayit_sayisi: int
    yinelenen_dosya: bool
    dosya_ozeti: str

    def sozluge_donustur(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class SurumGeriAlmaSonucu:
    hedef_surum_numarasi: int
    yeni_surum_numarasi: int
    etkinlestirilen_kayit_sayisi: int
    korunan_kayit_sayisi: int
    arsivlenen_kayit_sayisi: int

    def sozluge_donustur(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class DegisiklikPlani:
    korunan_kayitlar: tuple[tuple[Any, SSSKaydi], ...]
    guncellenen_kayitlar: tuple[tuple[Any, SSSKaydi], ...]
    eklenen_kayitlar: tuple[SSSKaydi, ...]
    arsivlenen_satirlar: tuple[Any, ...]


class ExcelDogrulamaHatasi(ValueError):
    pass


class ExcelYukleyici:
    standart_kolonlar = {
        "kategori": ("kategori",),
        "soru": ("soru",),
        "cevap": ("cevap",),
        "kaynak_dokuman": ("kaynak dokuman", "kaynak belge", "kaynak"),
    }
    uyds_kolonlar = {
        "modul": ("modul",),
        "musteri": ("musteri",),
        "madde_tarihi": ("madde tarihi",),
        "harici_kod": ("madde no",),
        "konu": ("konu",),
        "aciklama": ("aciklama",),
        "yapilacak": ("yapilacak",),
        "yapilan_islem": ("yapilan islem",),
    }

    def __init__(
        self,
        maksimum_dosya_boyutu: int | None = None,
        metin_parcalayici: MetinParcalayici | None = None,
    ) -> None:
        self.maksimum_dosya_boyutu = (
            maksimum_dosya_boyutu
            if maksimum_dosya_boyutu is not None
            else konfigurasyon.DOSYA_YUKLEME_MAKS_BOYUT
        )
        self.metin_parcalayici = metin_parcalayici or MetinParcalayici()

    def ice_aktar(
        self,
        dosya_yolu: str | Path,
        bildirilen_mime_tipi: str | None = None,
    ) -> IceAktarimSonucu:
        with _ICE_AKTARIM_KILIDI:
            guvenli_dosya_yolu = self.dosyayi_dogrula(
                dosya_yolu,
                bildirilen_mime_tipi,
            )
            dosya_ozeti = self._dosya_ozetini_hesapla(guvenli_dosya_yolu)
            tablolari_olustur()
            kayitlar = self._exceli_oku(guvenli_dosya_yolu)
            if not kayitlar:
                raise ExcelDogrulamaHatasi(
                    "Excel dosyasinda aktarilabilir kayit bulunamadi"
                )
            sonuc = self._kayitlari_aktar(
                guvenli_dosya_yolu.name,
                dosya_ozeti,
                kayitlar,
            )
        gunlukleyici.info(
            "Excel ice aktarimi tamamlandi: surum=%s, kayit=%s, eklenen=%s, guncellenen=%s, korunan=%s, arsivlenen=%s, parca=%s, yinelenen=%s",
            sonuc.surum_numarasi,
            sonuc.aktarilan_kayit_sayisi,
            sonuc.eklenen_kayit_sayisi,
            sonuc.guncellenen_kayit_sayisi,
            sonuc.korunan_kayit_sayisi,
            sonuc.arsivlenen_kayit_sayisi,
            sonuc.olusturulan_parca_sayisi,
            sonuc.yinelenen_dosya,
        )
        return sonuc

    def dosyayi_dogrula(
        self,
        dosya_yolu: str | Path,
        bildirilen_mime_tipi: str | None = None,
    ) -> Path:
        yol = Path(dosya_yolu).expanduser().resolve()
        if not yol.is_file():
            raise ExcelDogrulamaHatasi("Excel dosyasi bulunamadi")
        if yol.suffix.lower() != ".xlsx":
            raise ExcelDogrulamaHatasi("Yalnizca makrosuz .xlsx dosyalari kabul edilir")
        dosya_boyutu = yol.stat().st_size
        if dosya_boyutu <= 0:
            raise ExcelDogrulamaHatasi("Excel dosyasi bos")
        if dosya_boyutu > self.maksimum_dosya_boyutu:
            raise ExcelDogrulamaHatasi("Excel dosyasi izin verilen boyutu asiyor")
        mime_haritasi = mimetypes.MimeTypes()
        mime_haritasi.add_type(XLSX_MIME_TIPI, ".xlsx", strict=True)
        tahmin_edilen_mime_tipi = mime_haritasi.guess_type(yol.name)[0]
        if tahmin_edilen_mime_tipi != XLSX_MIME_TIPI:
            raise ExcelDogrulamaHatasi("Excel dosyasinin MIME tipi gecersiz")
        if bildirilen_mime_tipi:
            temiz_mime_tipi = bildirilen_mime_tipi.split(";", 1)[0].strip().lower()
            if temiz_mime_tipi not in IZINLI_MIME_TIPLERI:
                raise ExcelDogrulamaHatasi("Bildirilen MIME tipi gecersiz")
        with yol.open("rb") as dosya:
            if dosya.read(len(OOXML_IMZASI)) != OOXML_IMZASI:
                raise ExcelDogrulamaHatasi("Dosya gecerli bir OOXML arsivi degil")
        self._ooxml_arsivini_dogrula(yol)
        self._calisma_kitabini_dogrula(yol)
        return yol

    def _ooxml_arsivini_dogrula(self, dosya_yolu: Path) -> None:
        try:
            with zipfile.ZipFile(dosya_yolu) as arsiv:
                girisler = arsiv.infolist()
                if len(girisler) > AZAMI_ARSIV_GIRISI:
                    raise ExcelDogrulamaHatasi("Excel arsivinde cok fazla dosya bulunuyor")
                toplam_boyut = 0
                for giris in girisler:
                    temiz_ad = giris.filename.replace("\\", "/")
                    arsiv_yolu = PurePosixPath(temiz_ad)
                    if arsiv_yolu.is_absolute() or ".." in arsiv_yolu.parts:
                        raise ExcelDogrulamaHatasi("Excel arsivinde gecersiz dosya yolu var")
                    kucuk_ad = f"/{temiz_ad.lower().lstrip('/')}"
                    if any(parca in kucuk_ad for parca in TEHLIKELI_ARSIV_PARCALARI):
                        raise ExcelDogrulamaHatasi(
                            "Excel dosyasinda etkin veya harici icerik bulundu"
                        )
                    toplam_boyut += giris.file_size
                    if toplam_boyut > AZAMI_SIKISTIRILMAMIS_BOYUT:
                        raise ExcelDogrulamaHatasi(
                            "Excel arsivinin acilmis boyutu guvenlik sinirini asiyor"
                        )
                    if giris.compress_size > 0:
                        sikistirma_orani = giris.file_size / giris.compress_size
                        if sikistirma_orani > AZAMI_SIKISTIRMA_ORANI:
                            raise ExcelDogrulamaHatasi(
                                "Excel arsivinde supheli sikistirma orani bulundu"
                            )
                icerik_turleri = arsiv.read("[Content_Types].xml").lower()
                if b"spreadsheetml.sheet.main+xml" not in icerik_turleri:
                    raise ExcelDogrulamaHatasi(
                        "OOXML arsivi bir Excel calisma kitabi degil"
                    )
                if b"macroenabled" in icerik_turleri or b"vbaproject" in icerik_turleri:
                    raise ExcelDogrulamaHatasi("Makro iceren Excel dosyalari reddedilir")
        except zipfile.BadZipFile as hata:
            raise ExcelDogrulamaHatasi("Excel dosyasi bozuk veya gecersiz") from hata
        except KeyError as hata:
            raise ExcelDogrulamaHatasi("OOXML icerik tanimi bulunamadi") from hata

    def _calisma_kitabini_dogrula(self, dosya_yolu: Path) -> None:
        try:
            calisma_kitabi = load_workbook(
                dosya_yolu,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            try:
                if not calisma_kitabi.sheetnames:
                    raise ExcelDogrulamaHatasi("Excel dosyasinda calisma sayfasi yok")
                for sayfa in calisma_kitabi.worksheets:
                    for satir in sayfa.iter_rows():
                        for hucre in satir:
                            if hucre.data_type == "f":
                                raise ExcelDogrulamaHatasi(
                                    "Excel dosyasinda formul kullanilamaz"
                                )
                            if isinstance(hucre.value, str) and FORMUL_BASLANGICI.match(
                                hucre.value
                            ):
                                raise ExcelDogrulamaHatasi(
                                    "Excel dosyasinda formul enjeksiyonu riski bulundu"
                                )
            finally:
                calisma_kitabi.close()
        except ExcelDogrulamaHatasi:
            raise
        except Exception as hata:
            raise ExcelDogrulamaHatasi("Excel calisma kitabi okunamadi") from hata

    def _exceli_oku(self, dosya_yolu: Path) -> list[SSSKaydi]:
        baslik_satiri, esleme_turu = self._baslik_satirini_bul(dosya_yolu)
        kaynak_sayfa = self._ilk_sayfa_adini_al(dosya_yolu)
        veri_cercevesi = pd.read_excel(
            dosya_yolu,
            sheet_name=0,
            header=baslik_satiri,
            engine="openpyxl",
            dtype=object,
        )
        normalize_edilmis_kolonlar = {
            self._adi_normalize_et(kolon): kolon for kolon in veri_cercevesi.columns
        }
        if esleme_turu == "standart":
            kolon_eslemesi = self._kolonlari_esle(
                normalize_edilmis_kolonlar,
                self.standart_kolonlar,
            )
            return self._standart_kayitlari_olustur(
                veri_cercevesi,
                kolon_eslemesi,
                baslik_satiri,
                kaynak_sayfa,
            )
        kolon_eslemesi = self._kolonlari_esle(
            normalize_edilmis_kolonlar,
            self.uyds_kolonlar,
        )
        return self._uyds_kayitlarini_olustur(
            veri_cercevesi,
            kolon_eslemesi,
            baslik_satiri,
            kaynak_sayfa,
        )

    def _ilk_sayfa_adini_al(self, dosya_yolu: Path) -> str:
        calisma_kitabi = load_workbook(
            dosya_yolu,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            return calisma_kitabi.sheetnames[0]
        finally:
            calisma_kitabi.close()

    def _baslik_satirini_bul(self, dosya_yolu: Path) -> tuple[int, str]:
        onizleme = pd.read_excel(
            dosya_yolu,
            sheet_name=0,
            header=None,
            nrows=25,
            engine="openpyxl",
            dtype=object,
        )
        standart_zorunlu = {"kategori", "soru", "cevap", "kaynak dokuman"}
        uyds_zorunlu = {
            "modul",
            "madde tarihi",
            "madde no",
            "konu",
            "aciklama",
            "yapilacak",
            "yapilan islem",
        }
        for satir_numarasi, satir in onizleme.iterrows():
            basliklar = {
                self._adi_normalize_et(deger)
                for deger in satir.tolist()
                if not self._bos_mu(deger)
            }
            if standart_zorunlu.issubset(basliklar):
                return int(satir_numarasi), "standart"
            if uyds_zorunlu.issubset(basliklar):
                return int(satir_numarasi), "uyds"
        raise ExcelDogrulamaHatasi(
            "Zorunlu Excel sutunlari ilk 25 satirda bulunamadi"
        )

    def _kolonlari_esle(
        self,
        mevcut_kolonlar: dict[str, Any],
        beklenen_kolonlar: dict[str, tuple[str, ...]],
    ) -> dict[str, Any]:
        esleme: dict[str, Any] = {}
        for alan_adi, adaylar in beklenen_kolonlar.items():
            for aday in adaylar:
                if aday in mevcut_kolonlar:
                    esleme[alan_adi] = mevcut_kolonlar[aday]
                    break
        eksik_alanlar = set(beklenen_kolonlar) - set(esleme)
        if eksik_alanlar:
            raise ExcelDogrulamaHatasi(
                f"Eksik Excel alanlari: {', '.join(sorted(eksik_alanlar))}"
            )
        return esleme

    def _standart_kayitlari_olustur(
        self,
        veri_cercevesi: pd.DataFrame,
        kolonlar: dict[str, Any],
        baslik_satiri: int,
        kaynak_sayfa: str,
    ) -> list[SSSKaydi]:
        kayitlar: list[SSSKaydi] = []
        hatali_satirlar: list[int] = []
        gorulen_ozetler: set[str] = set()
        for indeks, satir in veri_cercevesi.iterrows():
            excel_satiri = int(indeks) + baslik_satiri + 2
            kategori = self._metne_donustur(satir[kolonlar["kategori"]])
            soru = self._metne_donustur(satir[kolonlar["soru"]])
            cevap = self._metne_donustur(satir[kolonlar["cevap"]])
            kaynak = self._metne_donustur(satir[kolonlar["kaynak_dokuman"]])
            if not any((kategori, soru, cevap, kaynak)):
                continue
            if not kategori or not soru or not cevap:
                hatali_satirlar.append(excel_satiri)
                continue
            kayit = self._kayit_olustur(
                kategori=kategori,
                soru=soru,
                cevap=cevap,
                kaynak_dokuman=kaynak,
                harici_kod=None,
                modul=kategori,
                kullanici_rolu=None,
                kaynak_sayfa=kaynak_sayfa,
                kaynak_satir=excel_satiri,
            )
            self._benzersiz_kaydi_ekle(kayitlar, gorulen_ozetler, kayit)
        self._hatali_satirlari_dogrula(hatali_satirlar)
        return kayitlar

    def _uyds_kayitlarini_olustur(
        self,
        veri_cercevesi: pd.DataFrame,
        kolonlar: dict[str, Any],
        baslik_satiri: int,
        kaynak_sayfa: str,
    ) -> list[SSSKaydi]:
        kayitlar: list[SSSKaydi] = []
        hatali_satirlar: list[int] = []
        gorulen_ozetler: set[str] = set()
        for indeks, satir in veri_cercevesi.iterrows():
            excel_satiri = int(indeks) + baslik_satiri + 2
            modul = self._metne_donustur(satir[kolonlar["modul"]])
            konu = self._metne_donustur(satir[kolonlar["konu"]])
            aciklama = self._metne_donustur(satir[kolonlar["aciklama"]])
            yapilacak = self._metne_donustur(satir[kolonlar["yapilacak"]])
            yapilan_islem = self._metne_donustur(
                satir[kolonlar["yapilan_islem"]]
            )
            harici_kod = self._metne_donustur(satir[kolonlar["harici_kod"]])
            madde_tarihi = self._metne_donustur(
                satir[kolonlar["madde_tarihi"]]
            )
            if not any(
                (
                    modul,
                    konu,
                    aciklama,
                    yapilacak,
                    yapilan_islem,
                    harici_kod,
                    madde_tarihi,
                )
            ):
                continue
            modul = modul or "Belirtilmemiş"
            soru = konu or aciklama or yapilacak or yapilan_islem
            cevap = yapilan_islem or yapilacak or aciklama
            if not soru or not cevap:
                hatali_satirlar.append(excel_satiri)
                continue
            kaynak_parcalari = ["UYDS SSS Listesi"]
            if harici_kod:
                kaynak_parcalari.append(f"Madde {harici_kod}")
            if madde_tarihi:
                kaynak_parcalari.append(madde_tarihi)
            kayit = self._kayit_olustur(
                kategori=modul,
                soru=soru,
                cevap=cevap,
                kaynak_dokuman=", ".join(kaynak_parcalari),
                harici_kod=harici_kod,
                modul=modul,
                kullanici_rolu=None,
                kaynak_sayfa=kaynak_sayfa,
                kaynak_satir=excel_satiri,
            )
            self._benzersiz_kaydi_ekle(kayitlar, gorulen_ozetler, kayit)
        self._hatali_satirlari_dogrula(hatali_satirlar)
        return kayitlar

    def _kayit_olustur(
        self,
        kategori: str,
        soru: str,
        cevap: str,
        kaynak_dokuman: str | None,
        harici_kod: str | None,
        modul: str | None,
        kullanici_rolu: str | None,
        kaynak_sayfa: str,
        kaynak_satir: int,
    ) -> SSSKaydi:
        arama_metni = self._metni_temizle(
            " ".join(
                parca
                for parca in (
                    kategori,
                    modul,
                    harici_kod,
                    soru,
                    cevap,
                    kaynak_dokuman,
                )
                if parca
            )
        )
        kayit_ozeti = self._kayit_ozeti_olustur(
            kategori,
            soru,
            cevap,
            kaynak_dokuman,
            harici_kod,
            modul,
            kullanici_rolu,
            arama_metni,
        )
        kayit_anahtari = self._kayit_anahtari_olustur(
            kategori,
            soru,
            harici_kod,
            modul,
        )
        return SSSKaydi(
            kategori=kategori,
            soru=soru,
            cevap=cevap,
            kaynak_dokuman=kaynak_dokuman,
            harici_kod=harici_kod,
            modul=modul,
            kullanici_rolu=kullanici_rolu,
            arama_metni=arama_metni,
            kaynak_sayfa=kaynak_sayfa,
            kaynak_satir=kaynak_satir,
            kayit_ozeti=kayit_ozeti,
            kayit_anahtari=kayit_anahtari,
        )

    def _kayit_ozeti_olustur(
        self,
        kategori: str,
        soru: str,
        cevap: str,
        kaynak_dokuman: str | None,
        harici_kod: str | None,
        modul: str | None,
        kullanici_rolu: str | None,
        arama_metni: str,
    ) -> str:
        ozet_verisi = {
            "arama_metni": arama_metni,
            "cevap": cevap,
            "harici_kod": harici_kod,
            "kategori": kategori,
            "kaynak_dokuman": kaynak_dokuman,
            "kullanici_rolu": kullanici_rolu,
            "modul": modul,
            "soru": soru,
        }
        return hashlib.sha256(
            json.dumps(
                ozet_verisi,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

    def _kayit_anahtari_olustur(
        self,
        kategori: str,
        soru: str,
        harici_kod: str | None,
        modul: str | None,
    ) -> str:
        anahtar_verisi = {
            "harici_kod": self._adi_normalize_et(harici_kod or ""),
            "kategori": self._adi_normalize_et(kategori),
            "modul": self._adi_normalize_et(modul or ""),
            "soru": self._adi_normalize_et(soru),
        }
        return hashlib.sha256(
            json.dumps(
                anahtar_verisi,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

    def _benzersiz_kaydi_ekle(
        self,
        kayitlar: list[SSSKaydi],
        gorulen_ozetler: set[str],
        kayit: SSSKaydi,
    ) -> None:
        if kayit.kayit_ozeti not in gorulen_ozetler:
            gorulen_ozetler.add(kayit.kayit_ozeti)
            kayitlar.append(kayit)

    def _hatali_satirlari_dogrula(self, hatali_satirlar: list[int]) -> None:
        if not hatali_satirlar:
            return
        gosterilenler = ", ".join(str(satir) for satir in hatali_satirlar[:20])
        devam_metni = " ve digerleri" if len(hatali_satirlar) > 20 else ""
        raise ExcelDogrulamaHatasi(
            f"Zorunlu alanlari eksik satirlar: {gosterilenler}{devam_metni}"
        )

    def kuru_calistir(
        self,
        dosya_yolu: str | Path,
        bildirilen_mime_tipi: str | None = None,
    ) -> KuruCalistirmaSonucu:
        with _ICE_AKTARIM_KILIDI:
            guvenli_dosya_yolu = self.dosyayi_dogrula(
                dosya_yolu,
                bildirilen_mime_tipi,
            )
            dosya_ozeti = self._dosya_ozetini_hesapla(guvenli_dosya_yolu)
            tablolari_olustur()
            kayitlar = self._exceli_oku(guvenli_dosya_yolu)
            if not kayitlar:
                raise ExcelDogrulamaHatasi(
                    "Excel dosyasinda analiz edilebilir kayit bulunamadi"
                )
            with baglam_yoneticisi() as baglanti:
                yinelenen_dosya = bool(
                    baglanti.execute(
                        """
                        SELECT 1
                        FROM veri_surumleri
                        WHERE dosya_ozeti = ?
                        """,
                        (dosya_ozeti,),
                    ).fetchone()
                )
                degisiklik_plani = self._degisiklik_planini_olustur(
                    baglanti,
                    kayitlar,
                )
        return KuruCalistirmaSonucu(
            toplam_kayit_sayisi=len(kayitlar),
            eklenen_kayit_sayisi=len(degisiklik_plani.eklenen_kayitlar),
            guncellenen_kayit_sayisi=len(
                degisiklik_plani.guncellenen_kayitlar
            ),
            korunan_kayit_sayisi=len(degisiklik_plani.korunan_kayitlar),
            arsivlenen_kayit_sayisi=len(
                degisiklik_plani.arsivlenen_satirlar
            ),
            yinelenen_dosya=yinelenen_dosya,
            dosya_ozeti=dosya_ozeti,
        )

    def surumu_geri_al(self, hedef_surum_numarasi: int) -> SurumGeriAlmaSonucu:
        if hedef_surum_numarasi < 1:
            raise ValueError("Hedef surum numarasi en az 1 olmalidir")
        with _ICE_AKTARIM_KILIDI:
            tablolari_olustur()
            with baglam_yoneticisi() as baglanti:
                baglanti.execute("BEGIN IMMEDIATE")
                hedef_surum = baglanti.execute(
                    """
                    SELECT id, surum_numarasi
                    FROM veri_surumleri
                    WHERE surum_numarasi = ?
                    """,
                    (hedef_surum_numarasi,),
                ).fetchone()
                if hedef_surum is None:
                    raise ValueError("Geri alinacak veri surumu bulunamadi")
                hedef_kayitlar = baglanti.execute(
                    """
                    SELECT
                        sss_id,
                        kaynak_sayfa,
                        kaynak_satir,
                        kayit_ozeti
                    FROM veri_surumu_kayitlari
                    WHERE veri_surumu_id = ?
                    ORDER BY sss_id
                    """,
                    (hedef_surum["id"],),
                ).fetchall()
                if not hedef_kayitlar:
                    raise ValueError("Hedef veri surumunde kayit bulunamadi")
                mevcut_aktif_kimlikler = {
                    satir["id"]
                    for satir in baglanti.execute(
                        """
                        SELECT id
                        FROM sss_kayitlari
                        WHERE durum = 'aktif'
                        """
                    ).fetchall()
                }
                hedef_kimlikler = {satir["sss_id"] for satir in hedef_kayitlar}
                onceki_surum = int(
                    baglanti.execute(
                        "SELECT COALESCE(MAX(surum_numarasi), 0) FROM veri_surumleri"
                    ).fetchone()[0]
                )
                yeni_surum = onceki_surum + 1
                korunan_kayit_sayisi = len(
                    mevcut_aktif_kimlikler & hedef_kimlikler
                )
                etkinlestirilen_kayit_sayisi = len(
                    hedef_kimlikler - mevcut_aktif_kimlikler
                )
                arsivlenen_kayit_sayisi = len(
                    mevcut_aktif_kimlikler - hedef_kimlikler
                )
                baglanti.execute(
                    """
                    UPDATE sss_kayitlari
                    SET aktif_mi = 0, durum = 'pasif'
                    WHERE durum = 'aktif'
                    """
                )
                baglanti.execute(
                    """
                    UPDATE veri_surumleri
                    SET durum = 'arsivlendi'
                    WHERE durum = 'aktif'
                    """
                )
                baglanti.execute(
                    """
                    INSERT INTO veri_surumleri (
                        dosya_adi,
                        surum_numarasi,
                        aktarilan_kayit_sayisi,
                        onceki_surum_numarasi,
                        durum,
                        eklenen_kayit_sayisi,
                        guncellenen_kayit_sayisi,
                        korunan_kayit_sayisi,
                        arsivlenen_kayit_sayisi,
                        olusturulan_parca_sayisi
                    )
                    VALUES (?, ?, ?, ?, 'aktif', ?, 0, ?, ?, 0)
                    """,
                    (
                        f"surum_{hedef_surum_numarasi}_geri_alma",
                        yeni_surum,
                        len(hedef_kayitlar),
                        onceki_surum or None,
                        etkinlestirilen_kayit_sayisi,
                        korunan_kayit_sayisi,
                        arsivlenen_kayit_sayisi,
                    ),
                )
                yeni_veri_surumu_id = baglanti.execute(
                    "SELECT last_insert_rowid()"
                ).fetchone()[0]
                baglanti.executemany(
                    """
                    UPDATE sss_kayitlari
                    SET
                        aktif_mi = 1,
                        durum = 'aktif',
                        son_gorulen_surum = ?,
                        aktarim_surumu_id = ?
                    WHERE id = ?
                    """,
                    [
                        (
                            yeni_surum,
                            yeni_veri_surumu_id,
                            satir["sss_id"],
                        )
                        for satir in hedef_kayitlar
                    ],
                )
                baglanti.executemany(
                    """
                    INSERT INTO veri_surumu_kayitlari (
                        veri_surumu_id,
                        sss_id,
                        kaynak_sayfa,
                        kaynak_satir,
                        kayit_ozeti
                    )
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            yeni_veri_surumu_id,
                            satir["sss_id"],
                            satir["kaynak_sayfa"],
                            satir["kaynak_satir"],
                            satir["kayit_ozeti"],
                        )
                        for satir in hedef_kayitlar
                    ],
                )
                denetim_detayi = json.dumps(
                    {
                        "hedef_surum_numarasi": hedef_surum_numarasi,
                        "yeni_surum_numarasi": yeni_surum,
                        "etkinlestirilen_kayit_sayisi": (
                            etkinlestirilen_kayit_sayisi
                        ),
                        "korunan_kayit_sayisi": korunan_kayit_sayisi,
                        "arsivlenen_kayit_sayisi": arsivlenen_kayit_sayisi,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                )
                baglanti.execute(
                    """
                    INSERT INTO sistem_gunlukleri (islem_turu, detay)
                    VALUES ('VERI_SURUMU_GERI_ALMA', ?)
                    """,
                    (denetim_detayi,),
                )
        return SurumGeriAlmaSonucu(
            hedef_surum_numarasi=hedef_surum_numarasi,
            yeni_surum_numarasi=yeni_surum,
            etkinlestirilen_kayit_sayisi=etkinlestirilen_kayit_sayisi,
            korunan_kayit_sayisi=korunan_kayit_sayisi,
            arsivlenen_kayit_sayisi=arsivlenen_kayit_sayisi,
        )

    def _degisiklik_planini_olustur(
        self,
        baglanti,
        kayitlar: list[SSSKaydi],
    ) -> DegisiklikPlani:
        mevcut_satirlar = [
            dict(satir)
            for satir in baglanti.execute(
                """
                SELECT
                    id,
                    kategori,
                    soru,
                    cevap,
                    kaynak_dokuman,
                    harici_kod,
                    modul,
                    kullanici_rolu,
                    arama_metni,
                    kayit_ozeti,
                    kayit_anahtari,
                    icerik_surumu,
                    revizyon_numarasi
                FROM sss_kayitlari
                WHERE durum = 'aktif'
                """
            ).fetchall()
        ]
        for mevcut_satir in mevcut_satirlar:
            if not mevcut_satir["kayit_ozeti"]:
                mevcut_satir["kayit_ozeti"] = self._kayit_ozeti_olustur(
                    mevcut_satir["kategori"],
                    mevcut_satir["soru"],
                    mevcut_satir["cevap"],
                    mevcut_satir["kaynak_dokuman"],
                    mevcut_satir["harici_kod"],
                    mevcut_satir["modul"],
                    mevcut_satir["kullanici_rolu"],
                    mevcut_satir["arama_metni"],
                )
            if not mevcut_satir["kayit_anahtari"]:
                mevcut_satir["kayit_anahtari"] = self._kayit_anahtari_olustur(
                    mevcut_satir["kategori"],
                    mevcut_satir["soru"],
                    mevcut_satir["harici_kod"],
                    mevcut_satir["modul"],
                )
        mevcut_ozet_haritasi = {
            satir["kayit_ozeti"]: satir for satir in mevcut_satirlar
        }
        korunan_kayitlar = tuple(
            (mevcut_ozet_haritasi[kayit.kayit_ozeti], kayit)
            for kayit in kayitlar
            if kayit.kayit_ozeti in mevcut_ozet_haritasi
        )
        korunan_mevcut_kimlikleri = {
            satir["id"] for satir, _ in korunan_kayitlar
        }
        korunan_yeni_ozetleri = {
            kayit.kayit_ozeti for _, kayit in korunan_kayitlar
        }
        kalan_mevcut_satirlar = [
            satir
            for satir in mevcut_satirlar
            if satir["id"] not in korunan_mevcut_kimlikleri
        ]
        kalan_yeni_kayitlar = [
            kayit
            for kayit in kayitlar
            if kayit.kayit_ozeti not in korunan_yeni_ozetleri
        ]
        mevcut_anahtar_gruplari = defaultdict(list)
        yeni_anahtar_gruplari = defaultdict(list)
        for satir in kalan_mevcut_satirlar:
            mevcut_anahtar_gruplari[satir["kayit_anahtari"]].append(satir)
        for kayit in kalan_yeni_kayitlar:
            yeni_anahtar_gruplari[kayit.kayit_anahtari].append(kayit)
        guncellenen_kayitlar = []
        guncellenen_mevcut_kimlikleri = set()
        guncellenen_yeni_ozetleri = set()
        for kayit in kalan_yeni_kayitlar:
            mevcut_grup = mevcut_anahtar_gruplari[kayit.kayit_anahtari]
            yeni_grup = yeni_anahtar_gruplari[kayit.kayit_anahtari]
            if len(mevcut_grup) != 1 or len(yeni_grup) != 1:
                continue
            mevcut_satir = mevcut_grup[0]
            guncellenen_kayitlar.append((mevcut_satir, kayit))
            guncellenen_mevcut_kimlikleri.add(mevcut_satir["id"])
            guncellenen_yeni_ozetleri.add(kayit.kayit_ozeti)
        eklenen_kayitlar = tuple(
            kayit
            for kayit in kalan_yeni_kayitlar
            if kayit.kayit_ozeti not in guncellenen_yeni_ozetleri
        )
        arsivlenen_satirlar = tuple(
            satir
            for satir in kalan_mevcut_satirlar
            if satir["id"] not in guncellenen_mevcut_kimlikleri
        )
        return DegisiklikPlani(
            korunan_kayitlar=korunan_kayitlar,
            guncellenen_kayitlar=tuple(guncellenen_kayitlar),
            eklenen_kayitlar=eklenen_kayitlar,
            arsivlenen_satirlar=arsivlenen_satirlar,
        )

    def _degisen_alanlari_bul(
        self,
        mevcut_satir: dict[str, Any],
        kayit: SSSKaydi,
    ) -> list[str]:
        alanlar = (
            "kategori",
            "soru",
            "cevap",
            "kaynak_dokuman",
            "harici_kod",
            "modul",
            "kullanici_rolu",
            "arama_metni",
        )
        return [
            alan
            for alan in alanlar
            if mevcut_satir.get(alan) != getattr(kayit, alan)
        ]

    def _kayitlari_aktar(
        self,
        dosya_adi: str,
        dosya_ozeti: str,
        kayitlar: list[SSSKaydi],
    ) -> IceAktarimSonucu:
        with baglam_yoneticisi() as baglanti:
            baglanti.execute("BEGIN IMMEDIATE")
            mevcut_surum = baglanti.execute(
                """
                SELECT
                    surum_numarasi,
                    aktarilan_kayit_sayisi,
                    durum
                FROM veri_surumleri
                WHERE dosya_ozeti = ?
                """,
                (dosya_ozeti,),
            ).fetchone()
            if mevcut_surum:
                if mevcut_surum["durum"] != "aktif":
                    raise ExcelDogrulamaHatasi(
                        "Dosya daha once arsivlenmis bir surum olarak aktarildi; "
                        "yeniden etkinlestirmek icin --surum-geri-al kullanin"
                    )
                return IceAktarimSonucu(
                    surum_numarasi=mevcut_surum["surum_numarasi"],
                    aktarilan_kayit_sayisi=mevcut_surum["aktarilan_kayit_sayisi"],
                    olusturulan_parca_sayisi=0,
                    arsivlenen_kayit_sayisi=0,
                    eklenen_kayit_sayisi=0,
                    guncellenen_kayit_sayisi=0,
                    korunan_kayit_sayisi=0,
                    yinelenen_dosya=True,
                    dosya_ozeti=dosya_ozeti,
                )
            onceki_surum = baglanti.execute(
                "SELECT COALESCE(MAX(surum_numarasi), 0) FROM veri_surumleri"
            ).fetchone()[0]
            yeni_surum = int(onceki_surum) + 1
            baglanti.execute(
                """
                INSERT INTO veri_surumleri (
                    dosya_adi,
                    surum_numarasi,
                    aktarilan_kayit_sayisi,
                    dosya_ozeti,
                    onceki_surum_numarasi,
                    durum
                )
                VALUES (?, ?, ?, ?, ?, 'aktif')
                """,
                (
                    dosya_adi,
                    yeni_surum,
                    len(kayitlar),
                    dosya_ozeti,
                    onceki_surum or None,
                ),
            )
            veri_surumu_id = baglanti.execute(
                "SELECT last_insert_rowid()"
            ).fetchone()[0]
            baglanti.execute(
                """
                UPDATE veri_surumleri
                SET durum = 'arsivlendi'
                WHERE durum = 'aktif' AND id != ?
                """,
                (veri_surumu_id,),
            )
            degisiklik_plani = self._degisiklik_planini_olustur(
                baglanti,
                kayitlar,
            )
            pasiflestirilecek_satirlar = [
                mevcut_satir
                for mevcut_satir, _ in degisiklik_plani.guncellenen_kayitlar
            ] + list(degisiklik_plani.arsivlenen_satirlar)
            if pasiflestirilecek_satirlar:
                baglanti.executemany(
                    """
                    UPDATE sss_kayitlari
                    SET aktif_mi = 0, durum = 'pasif'
                    WHERE id = ?
                    """,
                    [
                        (mevcut_satir["id"],)
                        for mevcut_satir in pasiflestirilecek_satirlar
                    ],
                )
            kayit_kimlikleri = {}
            for mevcut_satir, kayit in degisiklik_plani.korunan_kayitlar:
                baglanti.execute(
                    """
                    UPDATE sss_kayitlari
                    SET
                        kaynak_sayfa = ?,
                        kaynak_satir = ?,
                        kayit_ozeti = ?,
                        kayit_anahtari = ?,
                        son_gorulen_surum = ?,
                        aktarim_surumu_id = ?
                    WHERE id = ?
                    """,
                    (
                        kayit.kaynak_sayfa,
                        kayit.kaynak_satir,
                        kayit.kayit_ozeti,
                        kayit.kayit_anahtari,
                        yeni_surum,
                        veri_surumu_id,
                        mevcut_satir["id"],
                    ),
                )
                kayit_kimlikleri[kayit.kayit_ozeti] = mevcut_satir["id"]
            olusturulan_parca_sayisi = 0
            eklenecek_kayitlar = [
                (kayit, mevcut_satir, "guncelleme")
                for mevcut_satir, kayit in degisiklik_plani.guncellenen_kayitlar
            ] + [
                (kayit, None, "ekleme")
                for kayit in degisiklik_plani.eklenen_kayitlar
            ]
            eklenecek_harita = {
                kayit.kayit_ozeti: (kayit, mevcut_satir, degisiklik_turu)
                for kayit, mevcut_satir, degisiklik_turu in eklenecek_kayitlar
            }
            for kayit in kayitlar:
                if kayit.kayit_ozeti not in eklenecek_harita:
                    continue
                _, mevcut_satir, degisiklik_turu = eklenecek_harita[
                    kayit.kayit_ozeti
                ]
                onceki_kayit_id = (
                    mevcut_satir["id"] if mevcut_satir is not None else None
                )
                revizyon_numarasi = (
                    int(mevcut_satir["revizyon_numarasi"] or 1) + 1
                    if mevcut_satir is not None
                    else 1
                )
                icerik_surumu = (
                    int(mevcut_satir["icerik_surumu"] or 1) + 1
                    if mevcut_satir is not None
                    else 1
                )
                degisen_alanlar = (
                    json.dumps(
                        self._degisen_alanlari_bul(mevcut_satir, kayit),
                        ensure_ascii=False,
                    )
                    if mevcut_satir is not None
                    else None
                )
                imlec = baglanti.execute(
                    """
                    INSERT INTO sss_kayitlari (
                        kategori,
                        soru,
                        cevap,
                        kaynak_dokuman,
                        surum_numarasi,
                        aktif_mi,
                        harici_kod,
                        modul,
                        kullanici_rolu,
                        arama_metni,
                        durum,
                        icerik_surumu,
                        kaynak_sayfa,
                        kaynak_satir,
                        kayit_ozeti,
                        kayit_anahtari,
                        onceki_kayit_id,
                        revizyon_numarasi,
                        degisiklik_turu,
                        degisen_alanlar,
                        son_gorulen_surum,
                        aktarim_surumu_id
                    )
                    VALUES (
                        ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, 'aktif', ?,
                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                    )
                    """,
                    (
                        kayit.kategori,
                        kayit.soru,
                        kayit.cevap,
                        kayit.kaynak_dokuman,
                        yeni_surum,
                        kayit.harici_kod,
                        kayit.modul,
                        kayit.kullanici_rolu,
                        kayit.arama_metni,
                        icerik_surumu,
                        kayit.kaynak_sayfa,
                        kayit.kaynak_satir,
                        kayit.kayit_ozeti,
                        kayit.kayit_anahtari,
                        onceki_kayit_id,
                        revizyon_numarasi,
                        degisiklik_turu,
                        degisen_alanlar,
                        yeni_surum,
                        veri_surumu_id,
                    ),
                )
                sss_id = imlec.lastrowid
                kayit_kimlikleri[kayit.kayit_ozeti] = sss_id
                parca_metni = (
                    f"Kategori: {kayit.kategori}\n\n"
                    f"Soru: {kayit.soru}\n\n"
                    f"Cevap: {kayit.cevap}"
                )
                parcalar = self.metin_parcalayici.parcala(parca_metni)
                baglanti.executemany(
                    """
                    INSERT INTO sss_parcalari (sss_id, parca_metni, parca_sira_no)
                    VALUES (?, ?, ?)
                    """,
                    [
                        (sss_id, parca, sira_numarasi)
                        for sira_numarasi, parca in enumerate(parcalar, 1)
                    ],
                )
                olusturulan_parca_sayisi += len(parcalar)
            baglanti.executemany(
                """
                INSERT INTO veri_surumu_kayitlari (
                    veri_surumu_id,
                    sss_id,
                    kaynak_sayfa,
                    kaynak_satir,
                    kayit_ozeti
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                [
                    (
                        veri_surumu_id,
                        kayit_kimlikleri[kayit.kayit_ozeti],
                        kayit.kaynak_sayfa,
                        kayit.kaynak_satir,
                        kayit.kayit_ozeti,
                    )
                    for kayit in kayitlar
                ],
            )
            arsivlenen_kayit_sayisi = len(
                degisiklik_plani.arsivlenen_satirlar
            )
            eklenen_kayit_sayisi = len(degisiklik_plani.eklenen_kayitlar)
            guncellenen_kayit_sayisi = len(
                degisiklik_plani.guncellenen_kayitlar
            )
            korunan_kayit_sayisi = len(degisiklik_plani.korunan_kayitlar)
            baglanti.execute(
                """
                UPDATE veri_surumleri
                SET
                    eklenen_kayit_sayisi = ?,
                    guncellenen_kayit_sayisi = ?,
                    korunan_kayit_sayisi = ?,
                    arsivlenen_kayit_sayisi = ?,
                    olusturulan_parca_sayisi = ?
                WHERE id = ?
                """,
                (
                    eklenen_kayit_sayisi,
                    guncellenen_kayit_sayisi,
                    korunan_kayit_sayisi,
                    arsivlenen_kayit_sayisi,
                    olusturulan_parca_sayisi,
                    veri_surumu_id,
                ),
            )
            denetim_detayi = json.dumps(
                {
                    "surum_numarasi": yeni_surum,
                    "aktarilan_kayit_sayisi": len(kayitlar),
                    "olusturulan_parca_sayisi": olusturulan_parca_sayisi,
                    "arsivlenen_kayit_sayisi": arsivlenen_kayit_sayisi,
                    "eklenen_kayit_sayisi": eklenen_kayit_sayisi,
                    "guncellenen_kayit_sayisi": guncellenen_kayit_sayisi,
                    "korunan_kayit_sayisi": korunan_kayit_sayisi,
                    "dosya_ozeti": dosya_ozeti,
                },
                ensure_ascii=False,
                sort_keys=True,
            )
            baglanti.execute(
                """
                INSERT INTO sistem_gunlukleri (islem_turu, detay)
                VALUES ('EXCEL_ICE_AKTARIM', ?)
                """,
                (denetim_detayi,),
            )
        return IceAktarimSonucu(
            surum_numarasi=yeni_surum,
            aktarilan_kayit_sayisi=len(kayitlar),
            olusturulan_parca_sayisi=olusturulan_parca_sayisi,
            arsivlenen_kayit_sayisi=arsivlenen_kayit_sayisi,
            eklenen_kayit_sayisi=eklenen_kayit_sayisi,
            guncellenen_kayit_sayisi=guncellenen_kayit_sayisi,
            korunan_kayit_sayisi=korunan_kayit_sayisi,
            yinelenen_dosya=False,
            dosya_ozeti=dosya_ozeti,
        )

    def _dosya_ozetini_hesapla(self, dosya_yolu: Path) -> str:
        ozetleyici = hashlib.sha256()
        with dosya_yolu.open("rb") as dosya:
            for parca in iter(lambda: dosya.read(1024 * 1024), b""):
                ozetleyici.update(parca)
        return ozetleyici.hexdigest()

    def _adi_normalize_et(self, deger: Any) -> str:
        if deger is None:
            return ""
        metin = unicodedata.normalize("NFKD", str(deger).strip().casefold())
        aksansiz_metin = "".join(
            karakter for karakter in metin if not unicodedata.combining(karakter)
        )
        latin_metin = aksansiz_metin.replace("ı", "i")
        return re.sub(r"\s+", " ", latin_metin)

    def _metne_donustur(self, deger: Any) -> str | None:
        if self._bos_mu(deger):
            return None
        if isinstance(deger, pd.Timestamp):
            return deger.strftime("%d.%m.%Y")
        if isinstance(deger, float) and deger.is_integer():
            return str(int(deger))
        return self._metni_temizle(str(deger))

    def _metni_temizle(self, metin: str) -> str:
        satirlar = [
            re.sub(r"[ \t]+", " ", satir).strip()
            for satir in metin.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        ]
        return "\n".join(satirlar).strip()

    def _bos_mu(self, deger: Any) -> bool:
        if deger is None:
            return True
        try:
            return bool(pd.isna(deger))
        except (TypeError, ValueError):
            return False
