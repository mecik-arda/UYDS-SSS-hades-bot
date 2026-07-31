import hashlib
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from uygulama.cekirdek.gunlukleyici import gunlukleyici
from uygulama.cekirdek.konfigurasyon import konfigurasyon


@dataclass(frozen=True)
class BelgeIcerik:
    ham_metin: str
    basliklar: list[str]
    tablolar: list[list[list[str]]]
    madde_imleri: list[str]
    sayfa_sayisi: int
    kaynak_dosya: str
    son_degistirilme_tarihi: datetime
    sha256_ozeti: str
    yetki_seviyesi: int


class BelgeAyiklamaHatasi(RuntimeError):
    def __init__(self, mesaj: str, hata_kodu: str = "") -> None:
        super().__init__(mesaj)
        self.hata_kodu = hata_kodu


class BelgeAyiklayici:
    def belge_ayikla(self, dosya_yolu: str) -> BelgeIcerik:
        yol = Path(dosya_yolu)
        if not yol.is_file():
            raise BelgeAyiklamaHatasi(
                f"Dosya bulunamadi: {dosya_yolu}",
                "DOSYA_BULUNAMADI",
            )
        dosya_boyutu = yol.stat().st_size
        maks_boyut = konfigurasyon.BELGE_MAKS_BOYUT_MB * 1024 * 1024
        if dosya_boyutu > maks_boyut:
            raise BelgeAyiklamaHatasi(
                f"Dosya boyutu siniri asildi: {dosya_boyutu} bayt",
                "DOSYA_COK_BUYUK",
            )
        sha256_ozeti = hashlib.sha256(
            yol.read_bytes()
        ).hexdigest()
        uzanti = yol.suffix.lower()
        format_haritasi = {
            ".xlsx": self.excel_metin_ayikla,
            ".pdf": self.pdf_metin_ayikla,
            ".docx": self.docx_metin_ayikla,
            ".html": self.html_metin_ayikla,
            ".htm": self.html_metin_ayikla,
        }
        ayiklayici = format_haritasi.get(uzanti)
        if ayiklayici is None:
            raise BelgeAyiklamaHatasi(
                f"Desteklenmeyen dosya formati: {uzanti}",
                "FORMAT_TANINMADI",
            )
        son_degistirilme = datetime.fromtimestamp(
            yol.stat().st_mtime,
            tz=timezone.utc,
        )
        ham_metin, basliklar, tablolar, madde_imleri, sayfa_sayisi = (
            ayiklayici(dosya_yolu)
        )
        yetki_seviyesi = self.yetki_meta_verisi_oku(
            dosya_yolu,
            uzanti,
            ham_metin,
        )
        return BelgeIcerik(
            ham_metin=ham_metin,
            basliklar=basliklar,
            tablolar=tablolar,
            madde_imleri=madde_imleri,
            sayfa_sayisi=sayfa_sayisi,
            kaynak_dosya=dosya_yolu,
            son_degistirilme_tarihi=son_degistirilme,
            sha256_ozeti=sha256_ozeti,
            yetki_seviyesi=yetki_seviyesi,
        )

    def pdf_metin_ayikla(
        self,
        dosya_yolu: str,
    ) -> tuple[str, list[str], list[list[list[str]]], list[str], int]:
        try:
            import fitz
        except ImportError:
            raise BelgeAyiklamaHatasi(
                "PyMuPDF (fitz) yuklu degil",
                "PDF_KUTUPHANE_EKSIK",
            ) from None
        try:
            belge = fitz.open(dosya_yolu)
        except Exception as hata:
            raise BelgeAyiklamaHatasi(
                f"PDF acilamadi: {type(hata).__name__}",
                "PDF_BOZUK",
            ) from hata
        if belge.is_encrypted:
            belge.close()
            raise BelgeAyiklamaHatasi(
                "PDF parola korumali",
                "PDF_SIFRELI",
            )
        tum_metin = []
        tum_basliklar = []
        tum_tablolar = []
        tum_madde_imleri = []
        sayfa_sayisi = len(belge)
        for sayfa in belge:
            metin = sayfa.get_text("text")
            if not metin.strip():
                continue
            tum_metin.append(metin)
            satirlar = metin.splitlines()
            for satir in satirlar:
                temiz = satir.strip()
                if not temiz:
                    continue
                if re.match(r"^[\d.]+\s+\w", temiz):
                    tum_basliklar.append(temiz)
                elif re.match(r"^[\-\*•]\s", temiz):
                    tum_madde_imleri.append(
                        temiz.lstrip("-*• ").strip()
                    )
            sayfa_tablolari = self._pdf_tablolarini_ayikla(sayfa)
            tum_tablolar.extend(sayfa_tablolari)
        belge.close()
        birlestirilmis_metin = "\n".join(tum_metin)
        if not birlestirilmis_metin.strip():
            raise BelgeAyiklamaHatasi(
                "PDF icerigi bos",
                "METIN_BOS",
            )
        return (
            birlestirilmis_metin,
            tum_basliklar,
            tum_tablolar,
            tum_madde_imleri,
            sayfa_sayisi,
        )

    def _pdf_tablolarini_ayikla(
        self,
        sayfa,
    ) -> list[list[list[str]]]:
        motor_secimi = konfigurasyon.PDF_TABLO_MOTORU
        if motor_secimi == "ham_metin":
            return []
        try:
            if motor_secimi in ("auto", "fitz"):
                tablolar = sayfa.find_tables(
                    strategy="lines_strict",
                )
                if tablolar and tablolar.tables:
                    return [
                        [
                            [hucre.strip() for hucre in satir if hucre.strip()]
                            for satir in tablo.extract()
                            if any(hucre.strip() for hucre in satir)
                        ]
                        for tablo in tablolar.tables
                        if tablo.extract()
                    ]
            if motor_secimi == "pdfplumber":
                return self._pdfplumber_ile_tablo_ayikla(sayfa)
        except Exception as hata:
            gunlukleyici.warning(
                "PDF tablo ayiklama basarisiz: %s",
                type(hata).__name__,
            )
        return []

    def _pdfplumber_ile_tablo_ayikla(
        self,
        sayfa,
    ) -> list[list[list[str]]]:
        return []

    def docx_metin_ayikla(
        self,
        dosya_yolu: str,
    ) -> tuple[str, list[str], list[list[list[str]]], list[str], int]:
        try:
            from docx import Document
        except ImportError:
            raise BelgeAyiklamaHatasi(
                "python-docx yuklu degil",
                "DOCX_KUTUPHANE_EKSIK",
            ) from None
        try:
            belge = Document(dosya_yolu)
        except Exception as hata:
            raise BelgeAyiklamaHatasi(
                f"DOCX acilamadi: {type(hata).__name__}",
                "DOCX_BOZUK",
            ) from hata
        paragraflar = []
        basliklar = []
        madde_imleri = []
        for para in belge.paragraflar:
            metin = para.text.strip()
            if not metin:
                continue
            paragraflar.append(metin)
            if para.style.name.startswith("Heading"):
                basliklar.append(metin)
            if para.style.name.startswith("List"):
                madde_imleri.append(metin)
        tablolar = []
        for tablo in belge.tables:
            satirlar = []
            for satir in tablo.rows:
                hucreler = [
                    hucre.text.strip()
                    for hucre in satir.cells
                    if hucre.text.strip()
                ]
                if hucreler:
                    satirlar.append(hucreler)
            if satirlar:
                tablolar.append(satirlar)
        birlestirilmis = "\n".join(paragraflar)
        if not birlestirilmis.strip():
            raise BelgeAyiklamaHatasi(
                "DOCX icerigi bos",
                "METIN_BOS",
            )
        return (
            birlestirilmis,
            basliklar,
            tablolar,
            madde_imleri,
            1,
        )

    def html_metin_ayikla(
        self,
        kaynak: str,
    ) -> tuple[str, list[str], list[list[list[str]]], list[str], int]:
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise BelgeAyiklamaHatasi(
                "beautifulsoup4 yuklu degil",
                "HTML_KUTUPHANE_EKSIK",
            ) from None
        try:
            if Path(kaynak).is_file():
                with open(kaynak, encoding="utf-8") as dosya:
                    soup = BeautifulSoup(
                        dosya.read(),
                        "lxml",
                    )
            else:
                soup = BeautifulSoup(kaynak, "lxml")
        except Exception as hata:
            raise BelgeAyiklamaHatasi(
                f"HTML ayrıştırılamadı: {type(hata).__name__}",
                "HTML_BOZUK",
            ) from hata
        for eleman in soup(["script", "style", "nav", "footer"]):
            eleman.decompose()
        basliklar = []
        for seviye in range(1, 7):
            for etiket in soup.find_all(f"h{seviye}"):
                metin = etiket.get_text(strip=True)
                if metin:
                    basliklar.append(metin)
        madde_imleri = []
        for li in soup.find_all("li"):
            metin = li.get_text(strip=True)
            if metin:
                madde_imleri.append(metin)
        tablolar = []
        for tablo in soup.find_all("table"):
            satirlar = []
            for tr in tablo.find_all("tr"):
                hucreler = [
                    td.get_text(strip=True)
                    for td in tr.find_all(["td", "th"])
                    if td.get_text(strip=True)
                ]
                if hucreler:
                    satirlar.append(hucreler)
            if satirlar:
                tablolar.append(satirlar)
        ham_metin = soup.get_text(separator="\n", strip=True)
        if not ham_metin.strip():
            raise BelgeAyiklamaHatasi(
                "HTML icerigi bos",
                "METIN_BOS",
            )
        return (
            ham_metin,
            basliklar,
            tablolar,
            madde_imleri,
            1,
        )

    def excel_metin_ayikla(
        self,
        dosya_yolu: str,
    ) -> tuple[str, list[str], list[list[list[str]]], list[str], int]:
        try:
            import pandas as pd
        except ImportError:
            raise BelgeAyiklamaHatasi(
                "pandas yuklu degil",
                "EXCEL_KUTUPHANE_EKSIK",
            ) from None
        try:
            veri = pd.read_excel(
                dosya_yolu,
                sheet_name=None,
            )
        except Exception as hata:
            raise BelgeAyiklamaHatasi(
                f"Excel okunamadi: {type(hata).__name__}",
                "EXCEL_BOZUK",
            ) from hata
        tum_metin = []
        for sayfa_adi, cerceve in veri.items():
            cerceve = cerceve.dropna(how="all")
            tum_metin.append(
                cerceve.to_string(
                    index=False,
                    header=True,
                )
            )
        birlestirilmis = "\n".join(tum_metin)
        if not birlestirilmis.strip():
            raise BelgeAyiklamaHatasi(
                "Excel icerigi bos",
                "METIN_BOS",
            )
        return (
            birlestirilmis,
            list(veri.keys()),
            [],
            [],
            len(veri),
        )

    def yetki_meta_verisi_oku(
        self,
        dosya_yolu: str,
        dosya_formati: str,
        ham_metin: str = "",
    ) -> int:
        varsayilan = konfigurasyon.BELGE_VARSAYILAN_YETKI
        if dosya_formati == ".xlsx":
            return self._excel_yetki_oku(ham_metin, varsayilan)
        if dosya_formati == ".html" or dosya_formati == ".htm":
            return self._html_yetki_oku(dosya_yolu, varsayilan)
        if dosya_formati == ".docx":
            return self._docx_yetki_oku(dosya_yolu, varsayilan)
        return varsayilan

    def _excel_yetki_oku(
        self,
        ham_metin: str,
        varsayilan: int,
    ) -> int:
        for satir in ham_metin.splitlines()[:10]:
            eslesme = re.match(
                r"#\s*YETKI_SEVIYESI:\s*(\d+)",
                satir,
                re.IGNORECASE,
            )
            if eslesme:
                return int(eslesme.group(1))
        return varsayilan

    def _html_yetki_oku(
        self,
        dosya_yolu: str,
        varsayilan: int,
    ) -> int:
        try:
            from bs4 import BeautifulSoup
            with open(dosya_yolu, encoding="utf-8") as dosya:
                soup = BeautifulSoup(
                    dosya.read(),
                    "lxml",
                )
            meta = soup.find(
                "meta",
                attrs={"name": "permission-level"},
            )
            if meta and meta.get("content"):
                return int(meta["content"])
        except Exception:
            pass
        return varsayilan

    def _docx_yetki_oku(
        self,
        dosya_yolu: str,
        varsayilan: int,
    ) -> int:
        try:
            from docx import Document
            belge = Document(dosya_yolu)
            ozellikler = belge.core_properties
            if hasattr(ozellikler, "custom_properties"):
                deger = ozellikler.custom_properties.get(
                    "yetki_seviyesi"
                )
                if deger is not None:
                    return int(deger)
        except Exception:
            pass
        return varsayilan